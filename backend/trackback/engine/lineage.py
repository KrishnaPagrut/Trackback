"""The lineage tracer (SPEC.md Section 6.4): links every number people typed into spreadsheets, decks and certificates back to the books,
with nobody having drawn the link.

Extraction
- xlsx: formulas, not cached values. Every numeric or formula cell outside column A becomes an item; its label comes from column A (plus
  the column header on table-shaped sheets) and the sheet title is context. Total rows are skipped. Formulas are evaluated by a small
  whitelisted evaluator (cell refs, ranges, SUM, MAX, MIN, IF with comparisons, string literals); never eval of raw text.
- pptx: first text box is the slide title, the second the body; each paragraph is a line ("slide N, line M").
- pdf / htm: numbered lines ("line N"); compliance sentences become status claims.
- Numbers carry display precision: $48.1M is 48,100,000 +/- 50,000; $1,544K +/- 500; 2.95x +/- 0.005; 2.8% +/- 0.05. Years, list numbers,
  day-of-month after a month name, quarter and region tokens (Q1, R1) and section references are skipped.

Matching
- The metric library is evaluated at the named ledger snapshots; a file's own snapshot is the latest at or before its creation.
- Candidates are metrics within the item's tolerance with a compatible unit; a zero from an empty period is never evidence.
- Label score = 0.5 x overlap/metric tokens + 0.5 x Jaccard; hints (TTM, regions, quarter-end dates, quarter tokens) adjust it.
- Version score 1.0 at the file's own snapshot, 0.4 if only at an earlier one (then the item is `stale`).
- Accept if raw label >= 0.3 (0.15 with strong hints). Confidence = 0.35 + 0.3 raw + 0.15 hinted + 0.1 version + 0.2 margin (cap 0.99).
- Formula cells with no match are `derived`; hardcodes with no match try a small sum of 2 or 3 account totals (confidence 0.8);
  otherwise `unlinked` with near misses (within 5%).
"""
from __future__ import annotations

import itertools
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .books import QEND, QLABEL, Books, MetricLibrary

STOPWORDS = {"the", "a", "an", "of", "and", "for", "to", "in", "on", "at", "by", "as", "our", "we", "is", "with", "from", "plus", "vs", "x", "s", "k", "m", "per",
             "amount", "item", "total", "fiscal", "ended", "ending", "as", "under", "permitted"}
MONTHS = {"january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december",
          "jan", "feb", "mar", "apr", "jun", "jul", "aug", "sep", "sept", "oct", "nov", "dec"}
COMPLIANCE = re.compile(r"(within our .*covenant|in compliance with|remain(s)? in compliance|compliant with)", re.I)


def tokens(s: str) -> set[str]:
    out = set()
    for t in re.split(r"[^a-z0-9]+", (s or "").lower()):
        if not t or t in STOPWORDS:
            continue
        out.add({"adj": "adjusted", "ttm": "trailing", "rev": "revenue", "mgmt": "management"}.get(t, t))
    if "trailing" in out or "ttm" in (s or "").lower():
        out |= {"trailing", "four", "quarters"}
    return out


@dataclass
class Item:
    item_id: str
    file: str
    locator: str
    label: str
    context: str
    kind: str                 # hardcode | formula | typed | status
    value: Any
    unit: str | None          # usd | ratio | pct | status | None (unknown: workbook number)
    tolerance: float
    display: str | None = None
    refs: list[str] = field(default_factory=list)
    exact: bool = False
    header: str = ""


# ======================================================================================================================== numbers
NUM = re.compile(r"(?<![A-Za-z0-9.])(\(?-?\$?\s?\d[\d,]*(?:\.\d+)?(?:\s?[MKB]\b|x\b|%)?\)?)(?:\s+to\s+1\.00)?", re.I)


def parse_numbers(text: str) -> list[dict]:
    """Numbers in a line with their display precision. Skips years, list numbers, dates, Q1/R1 tokens and section references."""
    out = []
    for m in NUM.finditer(text):
        raw = m.group(1)
        start = m.start(1)
        before = text[:start].rstrip()
        prev_word = re.split(r"\s+", before)[-1].lower().strip("(,") if before else ""
        core = raw.strip("() ").replace(" ", "")
        if re.fullmatch(r"\d+\.", raw.strip()) and start == 0:        # "1." list number
            continue
        if start == 0 and re.match(r"\d+\.\s", text):                  # numbered line prefix
            continue
        if prev_word in MONTHS or prev_word.rstrip(",") in MONTHS:     # day of month after a month name
            continue
        if prev_word in ("section", "sec", "item", "items", "page", "slide", "line"):
            continue
        if start > 0 and text[start - 1].isalpha():                    # Q1, R1, G01
            continue
        if start > 0 and text[start - 1] in "'’`":                 # Q4'22: a year shorthand
            continue
        digits = re.sub(r"[^\d.]", "", core)
        if not digits or digits == ".":
            continue
        if re.fullmatch(r"(19|20)\d\d", digits) and "$" not in core:   # years
            continue
        val = float(digits)
        dec = len(digits.split(".")[1]) if "." in digits else 0
        suffix = core[-1].upper() if core[-1].upper() in "MKB%X" else ""
        unit, mult = None, 1.0
        if "$" in core or suffix in "MKB" and suffix:
            unit = "usd"
        if suffix == "M":
            mult = 1e6
        elif suffix == "K":
            mult = 1e3
        elif suffix == "B":
            mult = 1e9
        elif suffix == "%":
            unit = "pct"
        elif suffix == "X":
            unit = "ratio"
        ratio_to_one = m.group(0).lower().rstrip().endswith("to 1.00")
        if ratio_to_one:
            unit = "ratio"
        if re.match(r"\s*basis points", text[m.end(1):], re.I):
            unit = "bps"
        if raw.startswith("(") and raw.endswith(")"):
            val = -val
        v = val * mult
        tol = 0.5 * (10 ** -dec) * mult
        if re.match(r"\s*(billion|million|thousand)", text[m.end(1):], re.I):
            w = re.match(r"\s*(billion|million|thousand)", text[m.end(1):], re.I).group(1).lower()
            mult = {"billion": 1e9, "million": 1e6, "thousand": 1e3}[w]
            unit = "usd" if "$" in core else unit
            v = val * mult
            tol = 0.5 * (10 ** -dec) * mult
        if unit is None and "$" not in core and dec == 0 and val < 10 and not ratio_to_one:   # small bare integers are counts or list numbers
            continue
        out.append({"value": v, "unit": unit, "tolerance": max(tol, 1e-9), "display": m.group(0).strip(), "pos": start})
    return out


# ======================================================================================================================== xlsx
class FormulaError(Exception):
    pass


def _col(c: str) -> int:
    n = 0
    for ch in c:
        n = n * 26 + ord(ch) - 64
    return n


def _colname(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


class SheetEval:
    """Whitelisted evaluator: numbers, strings, cell refs, ranges, + - * /, comparisons, SUM, MAX, MIN, IF."""
    TOK = re.compile(r'\s*(?:(\d+(?:\.\d+)?)|("[^"]*")|([A-Z]+\$?\d+:[A-Z]+\$?\d+)|(\$?[A-Z]+\$?\d+)|(SUM|MAX|MIN|IF)\b|(<=|>=|<>|[-+*/(),<>=]))')

    def __init__(self, ws):
        self.ws = ws
        self.cache: dict[str, Any] = {}
        self.refs: dict[str, set[str]] = {}

    def value(self, addr: str) -> Any:
        addr = addr.replace("$", "")
        if addr in self.cache:
            return self.cache[addr]
        v = self.ws[addr].value
        if isinstance(v, str) and v.startswith("="):
            self.cache[addr] = None  # cycle guard
            refs: set[str] = set()
            saved = (getattr(self, "_t", None), getattr(self, "_i", 0), getattr(self, "_refs", None))
            try:
                v = self._eval(v[1:], refs)
            finally:
                self._t, self._i, self._refs = saved
            self.refs[addr] = refs
        self.cache[addr] = v
        return v

    def _eval(self, expr: str, refs: set[str]) -> Any:
        toks, pos = [], 0
        while pos < len(expr):
            m = self.TOK.match(expr, pos)
            if not m or m.end() == pos:
                if expr[pos:].strip() == "":
                    break
                raise FormulaError(f"unsupported formula text near: {expr[pos:pos + 12]!r}")
            toks.append(m.groups())
            pos = m.end()
        self._t, self._i, self._refs = toks, 0, refs
        v = self._cmp()
        if self._i != len(self._t):
            raise FormulaError("trailing tokens")
        return v

    def _peek(self):
        return self._t[self._i] if self._i < len(self._t) else (None,) * 6

    def _take(self):
        t = self._peek()
        self._i += 1
        return t

    def _cmp(self):
        a = self._add()
        op = self._peek()[5]
        if op in ("<=", ">=", "<>", "<", ">", "="):
            self._take()
            b = self._add()
            return {"<=": a <= b, ">=": a >= b, "<>": a != b, "<": a < b, ">": a > b, "=": a == b}[op]
        return a

    def _add(self):
        v = self._mul()
        while self._peek()[5] in ("+", "-"):
            op = self._take()[5]
            r = self._mul()
            v = v + r if op == "+" else v - r
        return v

    def _mul(self):
        v = self._unary()
        while self._peek()[5] in ("*", "/"):
            op = self._take()[5]
            r = self._unary()
            v = v * r if op == "*" else v / r
        return v

    def _unary(self):
        if self._peek()[5] == "-":
            self._take()
            return -self._unary()
        return self._atom()

    def _range(self, rng: str) -> list:
        a, b = rng.replace("$", "").split(":")
        ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
        cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
        vals = []
        for c in range(_col(ca), _col(cb) + 1):
            for r in range(int(ra), int(rb) + 1):
                addr = f"{_colname(c)}{r}"
                self._refs.add(addr)
                v = self.value(addr)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    vals.append(v)
        return vals

    def _args(self) -> list:
        if self._take()[5] != "(":
            raise FormulaError("expected (")
        args = []
        while True:
            t = self._peek()
            if t[2]:
                self._take()
                args.append(self._range(t[2]))
            else:
                args.append(self._cmp())
            sep = self._take()[5]
            if sep == ")":
                return args
            if sep != ",":
                raise FormulaError("expected , or )")

    def _atom(self):
        num, s, rng, ref, fn, op = self._take()
        if num:
            return float(num)
        if s:
            return s[1:-1]
        if ref:
            addr = ref.replace("$", "")
            self._refs.add(addr)
            v = self.value(addr)
            return v if v is not None else 0.0
        if fn:
            saved_refs = self._refs
            args = self._args()
            self._refs = saved_refs
            flat = [x for a in args for x in (a if isinstance(a, list) else [a])]
            if fn == "SUM":
                return float(sum(flat))
            if fn == "MAX":
                return float(max(flat))
            if fn == "MIN":
                return float(min(flat))
            if fn == "IF":
                return args[1] if args[0] else (args[2] if len(args) > 2 else False)
        if op == "(":
            v = self._cmp()
            if self._take()[5] != ")":
                raise FormulaError("expected )")
            return v
        raise FormulaError(f"unexpected token {op or num or s}")


def extract_xlsx(path: Path, prefix: str) -> list[Item]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    items: list[Item] = []
    for ws in wb.worksheets:
        ev = SheetEval(ws)
        header = {c.column_letter: str(c.value) for c in ws[1] if isinstance(c.value, str)}
        table = len(header) >= 3 and ws["A1"].value and not str(ws["A1"].value).startswith(("Tidewall", "Lyft"))
        title = " ".join(str(ws[f"A{r}"].value or "") for r in (1, 2)) if not table else ws.title
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                if cell.column_letter == "A" or cell.value is None:
                    continue
                if table and cell.row == 1:
                    continue
                label_a = str(ws[f"A{cell.row}"].value or "")
                if label_a.strip().lower() in ("total", "totals"):
                    continue
                raw = cell.value
                is_formula = isinstance(raw, str) and raw.startswith("=")
                if not is_formula and not isinstance(raw, (int, float)):
                    continue
                try:
                    v = ev.value(cell.coordinate)
                except (FormulaError, ZeroDivisionError, TypeError) as e:
                    v = None
                    label_a += f" [formula error: {e}]"
                label = f"{label_a} {header.get(cell.column_letter, '')}".strip() if table else label_a
                kind = "formula" if is_formula else "hardcode"
                unit = "status" if isinstance(v, str) else None
                items.append(Item(item_id=f"{prefix}.{ws.title}.{cell.coordinate}", file=path.name, locator=f"{ws.title}!{cell.coordinate}", label=label, context=f"{title} {ws.title}",
                                  kind="status" if unit == "status" else kind, value=v, unit=unit, tolerance=0.005, exact=True, header=header.get(cell.column_letter, "") if table else "",
                                  refs=[f"{prefix}.{ws.title}.{r}" for r in sorted(ev.refs.get(cell.coordinate, set()))] if is_formula else []))
    return items


# ======================================================================================================================== pptx, pdf, htm
def _line_items(prefix: str, file: str, locator: str, label_line: str, context: str, item_base: str) -> list[Item]:
    out: list[Item] = []
    if COMPLIANCE.search(label_line):
        return [Item(item_id=item_base, file=file, locator=locator, label=label_line, context=context, kind="status", value="Compliant", unit="status", tolerance=0)]
    nums = parse_numbers(label_line)
    for k, n in enumerate(nums):
        lab = label_line[:n["pos"]].strip(" :,-") if k == 0 else label_line[nums[k - 1]["pos"] + len(nums[k - 1]["display"]):n["pos"]].strip(" :,-()")
        if not lab or len(tokens(lab)) == 0:
            lab = label_line[:n["pos"]].strip(" :,-")
        iid = item_base if len(nums) == 1 else f"{item_base}.{'abcdefgh'[k]}"
        out.append(Item(item_id=iid, file=file, locator=locator, label=lab, context=f"{context} {label_line}", kind="typed", value=n["value"], unit=n["unit"],
                        tolerance=n["tolerance"], display=n["display"]))
    return out


def extract_pptx(path: Path, prefix: str) -> list[Item]:
    from pptx import Presentation
    prs = Presentation(str(path))
    items: list[Item] = []
    for si, slide in enumerate(prs.slides, start=1):
        boxes = [s for s in slide.shapes if s.has_text_frame]
        if not boxes:
            continue
        title = boxes[0].text_frame.text
        body = boxes[1].text_frame.paragraphs if len(boxes) > 1 else []
        for li, p in enumerate(body, start=1):
            text = "".join(r.text for r in p.runs)
            items += _line_items(prefix, path.name, f"slide {si}, line {li}", text, title, f"{prefix}.s{si}.l{li}")
    return items


def extract_pdf(path: Path, prefix: str) -> list[Item]:
    from pypdf import PdfReader
    text = "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
    items: list[Item] = []
    head = " ".join(text.splitlines()[:3])
    for ln in text.splitlines():
        m = re.match(r"\s*(\d+)\.\s+(.*)", ln)
        if not m:
            continue
        n, body = int(m.group(1)), m.group(2)
        items += _line_items(prefix, path.name, f"line {n}", body, head, f"{prefix}.l{n}")
    return items


def extract_html_lines(path: Path, prefix: str, wanted: list[str]) -> list[Item]:
    """Real filings: pick the sentences that contain each wanted phrase, then extract their numbers (Lyft's as-filed exhibit)."""
    import html as H
    raw = path.read_text(encoding="utf-8", errors="ignore")
    text = H.unescape(re.sub(r"<[^>]+>", " ", raw))
    text = re.sub(r"\s+", " ", text)
    sents = [x.strip(" •�") for x in re.split(r"(?<=[.;])\s+(?=[A-Z•�])|\s[•�]\s", text)]
    items: list[Item] = []
    for i, phrase in enumerate(wanted, start=1):
        s = next((x for x in sents if phrase.lower() in x.lower()), None)
        if s:
            items += _line_items(prefix, path.name, f"sentence {i}", s.strip()[:400], "", f"{prefix}.p{i}")
    return items


# ======================================================================================================================== matching
UNIT_OK = {None: {"usd", "ratio", "pct", "bps"}, "usd": {"usd"}, "ratio": {"ratio"}, "pct": {"pct"}, "status": {"status"}, "bps": {"bps"}}


def label_score(item_text: str, metric) -> float:
    it = tokens(item_text)
    best = 0.0
    for cand in [metric.label] + list(metric.synonyms):
        mt = tokens(cand)
        if not mt or not it:
            continue
        ov = len(it & mt)
        s = 0.5 * (ov / len(mt)) + 0.5 * (ov / len(it | mt))
        best = max(best, s)
    return best


def hints(item: Item, metric, file_quarter: str | None) -> tuple[float, bool]:
    from .. import config as C
    txt = f"{item.label} {item.context}".lower()
    h, strong = 0.0, False
    if re.search(r"\bttm\b|trailing|four (fiscal )?quarters", txt):
        strong = True
        if metric.id.startswith("ttm_"):
            h += 0.25
        elif metric.family in ("pl", "account") and metric.unit == "usd":
            h -= 0.15
    regions = [r for r, name in C.REGIONS.items() if name.lower() in txt or re.search(rf"\b{r.lower()}\b", txt)]
    if regions:
        strong = True
        if metric.region in regions:
            h += 0.3
        elif metric.region:
            h -= 0.3
    for q, end in QEND.items():
        md = end.split(",")[0].lower()        # "march 31"
        if md in txt:
            strong = True
            if metric.quarter == q and metric.family in ("balance", "derived") or (metric.quarter == q and metric.id.startswith("ttm_")):
                h += 0.2
            elif metric.quarter and metric.family in ("balance", "derived"):
                h -= 0.1
    qm = re.findall(r"\bq([1-4])\b", txt)
    if qm and metric.quarter and file_quarter:
        want = {q for q in QLABEL if q.endswith(f"Q{qm[0]}") and q <= file_quarter}
        want = {max(want)} if want else set()
        if metric.quarter in want:
            h += 0.2
        else:
            h -= 0.1
    return h, strong


def _is_empty_zero(metric, value) -> bool:
    return isinstance(value, (int, float)) and abs(value) < 0.005 and metric.family in ("pl", "account", "balance")


def match(items: list[Item], lib: MetricLibrary, snapshot_values: dict[str, dict], file_snapshot: dict[str, str], snapshot_order: list[str],
          file_quarter: dict[str, str], account_metrics_by_q: dict[str, list[str]] | None = None, snapshot_labels: dict[str, str] | None = None) -> list[dict]:
    """snapshot_values: name -> metric values; file_snapshot: file -> snapshot name; snapshot_order: oldest first."""
    links: list[dict] = []
    by_id = {it.item_id: it for it in items}
    for it in items:
        own = file_snapshot[it.file]
        earlier = snapshot_order[:snapshot_order.index(own)]
        fq = file_quarter.get(it.file)
        cands = []
        head = tokens(it.header) - {"q1", "q2", "q3", "q4"}
        for mid, metric in lib.metrics.items():
            if metric.unit not in UNIT_OK.get(it.unit, set()) and not (it.unit == "status" and metric.unit == "status"):
                continue
            if head and not head & set().union(*(tokens(x) for x in [metric.label] + list(metric.synonyms))):
                continue          # table sheets: the column header names the quantity, so it must agree with the metric
            for snap, vsc in [(own, 1.0)] + [(s, 0.4) for s in reversed(earlier)]:
                v = snapshot_values[snap].get(mid)
                if v is None:
                    continue
                if it.unit == "status":
                    ok = isinstance(v, str) and v == it.value
                else:
                    if not isinstance(v, (int, float)) or not isinstance(it.value, (int, float)):
                        continue
                    if _is_empty_zero(metric, v):
                        continue
                    ok = abs(v - it.value) <= it.tolerance + 1e-9
                if ok:
                    raw = label_score(f"{it.label} {it.context}", metric)
                    h, strong = hints(it, metric, fq)
                    cands.append({"metric": mid, "snapshot": snap, "version": vsc, "raw": raw, "hinted": max(0.0, min(1.0, raw + h)), "strong": strong, "value": v})
                    break
        cands.sort(key=lambda c: (c["hinted"] + 0.3 * c["version"], c["raw"]), reverse=True)
        link = {"item_id": it.item_id, "file": it.file, "locator": it.locator, "label": it.label, "value": it.value, "unit": it.unit, "kind": it.kind,
                "display": it.display, "refs": it.refs, "tolerance": it.tolerance, "exact": it.exact, "context": (it.context or "").strip(), "candidates": [{k: c[k] for k in ("metric", "snapshot", "raw", "hinted", "version")} for c in cands[:4]]}
        best = cands[0] if cands else None
        if best and (best["raw"] >= 0.3 or (best["strong"] and best["raw"] >= 0.15)):
            runner = cands[1]["hinted"] + 0.3 * cands[1]["version"] if len(cands) > 1 else 0.0
            margin = max(0.0, min(1.0, best["hinted"] + 0.3 * best["version"] - runner))
            conf = min(0.99, 0.35 + 0.3 * best["raw"] + 0.15 * best["hinted"] + 0.1 * best["version"] + 0.2 * margin)
            link.update(metric_id=best["metric"], snapshot=best["snapshot"], confidence=round(conf, 3), method="matched")
            metric = lib.metrics[best["metric"]]
            if metric.family == "policy":
                link.update(status="policy", explanation=f"{metric.label}: a contract or policy term")
            elif best["snapshot"] != own:
                cur = snapshot_values[own].get(best["metric"])
                lab = (snapshot_labels or {}).get(best["snapshot"], f"{best['snapshot']} snapshot")
                link.update(status="stale", explanation=f"typed from the {lab} ({_fmt(best['value'], metric.unit)}); the books already said {_fmt(cur, metric.unit)} when the file was made")
            else:
                link.update(status="linked", explanation=f"{metric.label} = {_fmt(best['value'], metric.unit)} in the {(snapshot_labels or {}).get(own, own + ' snapshot')}")
        elif it.kind == "formula" and it.refs:
            link.update(status="derived", metric_id=None, snapshot=own, confidence=0.9, method="formula", explanation=f"formula over {', '.join(r.split('.')[-1] for r in it.refs)}")
        elif isinstance(it.value, (int, float)) and account_metrics_by_q:
            order = {q: account_metrics_by_q[q] for q in sorted(account_metrics_by_q, key=lambda q: (q != fq, q))}
            rec = _small_sum(it, snapshot_values[own], order, lib)
            if rec:
                link.update(status="reconstructed", metric_id="sum:" + "+".join(rec), snapshot=own, confidence=0.8, method="small_sum",
                            explanation=f"no single metric; reconstructed as {' + '.join(lib.metrics[m].label for m in rec)}")
            else:
                link.update(status="unlinked", metric_id=None, snapshot=own, confidence=0.0, method=None, explanation="no metric or small sum matches", near_misses=_near(it, lib, snapshot_values[own]))
        else:
            link.update(status="unlinked", metric_id=None, snapshot=own, confidence=0.0, method=None, explanation="no metric matches", near_misses=_near(it, lib, snapshot_values[own]))
        links.append(link)
    # derived formula cells that DID match a metric keep their refs; nothing else to do
    _ = by_id
    return links


def _small_sum(it: Item, values: dict, acct_by_q: dict[str, list[str]], lib: MetricLibrary) -> list[str] | None:
    """2 or 3 positive account totals of one quarter that hit the value within tolerance; among several, the one whose labels best
    explain the item's label (and at least a little)."""
    item_t = tokens(f"{it.label} {it.context}")
    for q, mids in acct_by_q.items():
        pos = [(m, values[m]) for m in mids if isinstance(values.get(m), (int, float)) and values[m] > 0]
        hits = []
        for k in (2, 3):
            for combo in itertools.combinations(pos, k):
                if abs(sum(v for _m, v in combo) - it.value) <= it.tolerance:
                    mt = set().union(*(tokens(lib.metrics[m].label) | set().union(*(tokens(x) for x in lib.metrics[m].synonyms)) for m, _v in combo))
                    hits.append((len(item_t & mt) / max(1, len(item_t)), [m for m, _v in combo]))
        hits = [h for h in hits if h[0] >= 0.15]
        if hits:
            return max(hits, key=lambda h: h[0])[1]
    return None


def _near(it: Item, lib: MetricLibrary, values: dict) -> list[dict]:
    out = []
    if not isinstance(it.value, (int, float)) or it.value == 0:
        return out
    for mid, v in values.items():
        if isinstance(v, (int, float)) and v and abs(v - it.value) / abs(it.value) <= 0.05:
            out.append({"metric": mid, "value": v, "diff_pct": round((v - it.value) / abs(it.value) * 100, 2)})
        # same digits at another power of ten: a typo signature
        if isinstance(v, (int, float)) and v and it.value and abs(v) > 0:
            for p in (10, 100, 1000, 0.1, 0.01, 0.001):
                if abs(v * p - it.value) <= max(it.tolerance, abs(it.value) * 1e-9):
                    out.append({"metric": mid, "value": v, "power_of_ten": p})
    return sorted(out, key=lambda x: abs(x.get("diff_pct", 0)))[:5]


def _fmt(v, unit) -> str:
    if v is None:
        return "n/a"
    if isinstance(v, str):
        return v
    if unit == "ratio":
        return f"{v:.2f}x"
    if unit == "pct":
        return f"{v:.1f}%"
    return f"${v:,.0f}"


def snapshot_values(lib: MetricLibrary, entries: list[dict], snapshots: dict[str, str]) -> dict[str, dict]:
    return {name: lib.evaluate(Books(entries, t)) for name, t in snapshots.items()}
