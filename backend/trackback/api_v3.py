"""The v3 API (SPEC.md Section 13), for both worlds (?world=tidewall|lyft). The UI reads every number from here, never from text in the UI.

/api/state  /api/stage  /api/review/run  /api/recons  /api/exceptions  /api/evaluations/{id}  /api/lineage  /api/files/{source}
/api/ripple/{id}  /api/materiality  /api/approvals  /api/approvals/{id}/decision  /api/people/{person}/act  /api/clock/advance
/api/memo/{id}  /api/memento/timemap/{source}  /api/memento/timegate/{source}  /api/replay/attempt  /api/ledger  /api/scorecard
/api/lessons  /api/lessons/{id}/approve  /api/corrections  /api/pings  /api/built/{artifact_id}  /api/events  /api/s1  /api/reset
"""
from __future__ import annotations

import json
import re
import shutil
import sqlite3
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException
from fastapi.responses import JSONResponse
from pydantic import BaseModel

from . import config as C
from .core import approvals as AP, db, graph as G
from .core.llm import gateway

router = APIRouter()
UNDO_DIR = C.DATA_DIR / "snapshots" / "undo"


def _latest(world: str) -> str | None:
    r = db.q1("SELECT value FROM kv WHERE key=?", (f"review.latest.{world}",))
    return r["value"] if r else None


def _kv(key: str, default=None):
    r = db.q1("SELECT value FROM kv WHERE key=?", (key,))
    return json.loads(r["value"]) if r else default


def _eval_file() -> dict:
    p = C.DATA_DIR / "eval" / "latest.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def _snapshot_for_undo() -> None:
    UNDO_DIR.mkdir(parents=True, exist_ok=True)
    n = len(list(UNDO_DIR.glob("*.sqlite")))
    src = sqlite3.connect(str(C.DB_PATH)); d = sqlite3.connect(str(UNDO_DIR / f"undo_{n:03d}.sqlite")); src.backup(d); d.close(); src.close()


# ------------------------------------------------------------------------------------------------------------------ basics
@router.get("/health")
def health():
    return {"ok": True, "now": C.KEY_DATES["now"], "mode": gateway.mode()}


@router.get("/state")
def state(world: str = "tidewall"):
    from . import workflow
    rid = _latest(world)
    ev = _eval_file()
    pend = db.q1("SELECT COUNT(*) n FROM actions WHERE world=? AND state='proposed'", (world,))["n"]
    held = db.q1("SELECT COUNT(*) n FROM holds WHERE world=? AND state='held'", (world,))["n"]
    cache = {r["who"]: r["n"] for r in db.q("SELECT CASE WHEN model LIKE 'scripted:%' THEN 'scripted' WHEN model LIKE 'fallback:%' THEN 'fallback' ELSE 'real' END who, COUNT(*) n FROM llm_cache GROUP BY 1")}
    base = {"world": world, "mode": gateway.mode(), "run_id": rid, "pending": pend, "holds": held, "headline": ev.get("headline"), "cache_rows": cache,
            "agent_model": C.AGENT_MODEL, "judge_model": C.JUDGE_MODEL, "undo_depth": len(list(UNDO_DIR.glob("*.sqlite"))) if UNDO_DIR.exists() else 0}
    if world == "lyft":
        from .world import lyft as LY
        return {**base, "company": {"name": "Lyft, Inc.", "real": True, "note": "Release as filed and the correction are real; the finance team's documents, agents and outside dependents are reconstructions."},
                "now": LY.T["review"], "clock": LY.T["review"], "people": LY.PEOPLE, "tape": LY.TAPE}
    return {**base, "company": C.COMPANY_V3, "now": C.KEY_DATES["now"], "clock": workflow.clock(), "people": C.PEOPLE_V3, "regions": C.REGIONS, "key_dates": C.KEY_DATES,
            "materiality": C.MATERIALITY_V3, "covenant": C.COVENANT}


# ------------------------------------------------------------------------------------------------------------------ the review
@router.post("/review/run")
def review_run():
    from . import review
    _snapshot_for_undo()
    out = review.run_all()
    return {"run_id": out["run_id"], "summary": out["summary"]}


@router.get("/recons")
def recons(world: str = "tidewall", recon: str | None = None, outcome: str | None = None, limit: int = 400):
    rid = _latest(world)
    if world == "lyft":
        from .world import lyft as LY
        links = [dict(r) for r in db.q("SELECT i.*, l.status, l.metric_id, l.confidence, l.explanation FROM trace_items i JOIN trace_links l ON l.item_id=i.item_id AND l.run_id=i.run_id WHERE i.run_id=?", (rid,))]
        summary = [{"recon": "release_to_sources", "title": "Release figures vs the close pack and the guidance model", "counts": {"matched": sum(1 for l in links if l["status"] in ("linked", "policy")), "cleared": 0,
                    "exceptions": sum(1 for l in links if l["status"] == "unlinked")}, "cleared_by_reason": {}}]
        return {"summary": summary, "items": [{"recon": "release_to_sources", "outcome": "matched" if l["status"] in ("linked", "policy") else "exception", "reason": l["status"], "amount": l["value_num"],
                                                "detail": {"label": l["label"], "display": l["display"], "metric": l["metric_id"], "explanation": l["explanation"]}} for l in links], "_": LY.T["review"]}
    summary = _kv(f"recon.{rid}", [])
    sql = "SELECT recon, outcome, reason, amount, detail_json, exception_id FROM recon_items WHERE run_id=?"
    params: list[Any] = [rid]
    if recon:
        sql += " AND recon=?"; params.append(recon)
    if outcome:
        sql += " AND outcome=?"; params.append(outcome)
    sql += " ORDER BY CASE outcome WHEN 'exception' THEN 0 WHEN 'cleared' THEN 1 ELSE 2 END, recon LIMIT ?"
    params.append(limit)
    items = [{**{k: r[k] for k in ("recon", "outcome", "reason", "amount", "exception_id")}, "detail": json.loads(r["detail_json"])} for r in db.q(sql, params)]
    return {"summary": summary, "items": items}


@router.get("/exceptions")
def exceptions(world: str = "tidewall"):
    rid = _latest(world)
    out = []
    for r in db.q("SELECT * FROM exceptions WHERE run_id=? ORDER BY exception_id", (rid,)):
        d = json.loads(r["detail_json"])
        fin = db.q1("SELECT claims_json FROM evaluations WHERE run_id=? AND exception_id=? AND playbook_version LIKE '%+jade'", (rid, r["exception_id"]))
        f = json.loads(fin["claims_json"]) if fin else {}
        rp = db.q1("SELECT counts_json FROM ripples WHERE run_id=? AND exception_id=?", (rid, r["exception_id"]))
        out.append({"exception_id": r["exception_id"], "recon": r["recon"], "kind": r["kind"], "title": r["title"], "amount": r["amount"], "system_of_record": r["system_of_record"],
                    "classification": (f.get("q3_classification") or {}).get("value"), "material": (f.get("q5_matters") or {}).get("material"), "treatment": (f.get("q7_fix") or {}).get("treatment"),
                    "counts": json.loads(rp["counts_json"]) if rp else None, "trust": json.loads(r["trust_why"] or "null"), "knowable_at": r["knowable_at"], "evidence": json.loads(r["evidence_json"] or "[]"),
                    "by_region": d.get("by_region")})
    return {"run_id": rid, "exceptions": out}


@router.get("/evaluations/{exception_id}")
def evaluation(exception_id: str, world: str = "tidewall"):
    rid = _latest(world)
    v1 = db.q1("SELECT claims_json, playbook_version FROM evaluations WHERE run_id=? AND exception_id=? AND playbook_version NOT LIKE '%+jade'", (rid, exception_id))
    fin = db.q1("SELECT claims_json FROM evaluations WHERE run_id=? AND exception_id=? AND playbook_version LIKE '%+jade'", (rid, exception_id))
    if v1 is None:
        raise HTTPException(404, "no evaluation")
    jade = [dict(r) | {"agent_value": json.loads(r["agent_value_json"] or "null"), "corrected": json.loads(r["corrected_json"] or "null")}
            for r in db.q("SELECT claim_key, verdict, why, method, agent_value_json, corrected_json FROM jade_reviews WHERE eval_id=? ORDER BY rowid", (f"EV-{rid}-{exception_id}",))]
    for j in jade:
        j.pop("agent_value_json", None); j.pop("corrected_json", None)
    return {"exception_id": exception_id, "playbook": v1["playbook_version"], "agent": json.loads(v1["claims_json"]), "jade": jade, "final": json.loads(fin["claims_json"]) if fin else None}


@router.get("/materiality")
def materiality(world: str = "tidewall"):
    rid = _latest(world)
    return _kv(f"materiality.{rid}", {}) | {"memo": C.MATERIALITY_V3}


# ------------------------------------------------------------------------------------------------------------------ lineage and files
@router.get("/lineage")
def lineage(world: str = "tidewall"):
    rid = _latest(world)
    rows = db.q("SELECT i.*, l.status, l.metric_id, l.snapshot, l.method, l.confidence, l.explanation, l.candidates_json FROM trace_items i JOIN trace_links l ON l.item_id=i.item_id AND l.run_id=i.run_id WHERE i.run_id=? ORDER BY i.file, i.rowid", (rid,))
    files: dict[str, list] = {}
    for r in rows:
        d = dict(r)
        d["candidates"] = json.loads(d.pop("candidates_json") or "[]")
        d["refs"] = json.loads(d.pop("refs_json") or "[]")
        files.setdefault(r["file"], []).append(d)
    return {"run_id": rid, "files": files, "counts": {s: sum(1 for r in rows if r["status"] == s) for s in ("linked", "policy", "stale", "derived", "reconstructed", "unlinked")}}


def _file_source(source: str) -> tuple[Path, dict]:
    v = db.q1("SELECT * FROM source_versions WHERE source_id=? ORDER BY tx_from DESC LIMIT 1", (source,))
    if v is None:
        raise HTTPException(404, "no such source")
    from .engine import replay as R
    ver = {"version_id": v["version_id"], "captured_at": v["tx_from"], "body": json.loads(v["body"])}
    return R.file_path(ver), ver


@router.get("/files/{source}")
def file_view(source: str, world: str = "tidewall", batch: str | None = None, period: str | None = None, q: str | None = None, limit: int = 60, version: str | None = None):
    """The rendered file with highlights: workbook cells (formulas evaluated by Trackback), deck lines, certificate lines, CSV rows, exhibit sentences."""
    from .engine import lineage as L, replay as R
    if version:
        row = db.q1("SELECT * FROM source_versions WHERE version_id=?", (version,))
        ver = {"version_id": row["version_id"], "captured_at": row["tx_from"], "body": json.loads(row["body"])}
        path = R.file_path(ver)
    else:
        path, ver = _file_source(source)
    rid = _latest(world)
    links = {r["item_id"]: dict(r) for r in db.q("SELECT i.item_id, i.locator, l.status, l.metric_id, l.explanation, l.confidence FROM trace_items i JOIN trace_links l ON l.item_id=i.item_id AND l.run_id=i.run_id WHERE i.run_id=? AND i.file=?", (rid, path.name))}
    by_loc = {v["locator"]: v for v in links.values()}
    title = db.q1("SELECT title FROM sources WHERE source_id=?", (source,))["title"]
    base = {"source": source, "title": title, "version_id": ver["version_id"], "captured_at": ver["captured_at"], "file": path.name}
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False)
        sheets = []
        for ws in wb.worksheets:
            ev = L.SheetEval(ws)
            rows = []
            for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
                cells = []
                for c in r:
                    if c.value is None:
                        cells.append(None); continue
                    raw = c.value
                    val = ev.value(c.coordinate) if isinstance(raw, str) and raw.startswith("=") else raw
                    lk = by_loc.get(f"{ws.title}!{c.coordinate}")
                    cells.append({"addr": c.coordinate, "value": val, "formula": raw if isinstance(raw, str) and raw.startswith("=") else None, "link": lk})
                rows.append(cells)
            sheets.append({"name": ws.title, "rows": rows})
        return {**base, "kind": "xlsx", "sheets": sheets}
    if path.suffix == ".pptx":
        from pptx import Presentation
        slides = []
        for si, s in enumerate(Presentation(str(path)).slides, start=1):
            boxes = [b for b in s.shapes if b.has_text_frame]
            lines = [{"n": li, "text": "".join(r.text for r in p.runs), "link": by_loc.get(f"slide {si}, line {li}"),
                      "links": [v for v in links.values() if v["locator"] == f"slide {si}, line {li}"]} for li, p in enumerate(boxes[1].text_frame.paragraphs if len(boxes) > 1 else [], start=1)]
            slides.append({"n": si, "title": boxes[0].text_frame.text if boxes else "", "lines": lines})
        return {**base, "kind": "pptx", "slides": slides}
    if path.suffix == ".pdf":
        from pypdf import PdfReader
        text = "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
        lines = []
        for ln in text.splitlines():
            m = re.match(r"\s*(\d+)\.\s", ln)
            lines.append({"text": ln, "link": by_loc.get(f"line {m.group(1)}") if m else None})
        return {**base, "kind": "pdf", "lines": lines}
    if path.suffix == ".csv":
        rows = R.rows(source, "9999-01-01T00:00:00+00:00") if not version else [dict(x) for x in R._read_file(str(path), path.stat().st_mtime)]
        if batch:
            rows = [r for r in rows if r.get("batch") == batch]
        if period:
            rows = [r for r in rows if r.get("period") == period]
        if q:
            rows = [r for r in rows if q.lower() in json.dumps(r).lower()]
        agg = None
        if source == "billing_release_lines":
            agg = {}
            for r in rows:
                k = f"{r['period']} {r['batch']}"
                agg[k] = round(agg.get(k, 0) + float(r["release_amount"]), 2)
        return {**base, "kind": "csv", "n": len(rows), "columns": list(rows[0].keys()) if rows else [], "rows": rows[:limit], "totals": agg}
    if path.suffix in (".htm", ".html"):
        rid_l = _latest("lyft")
        items = [dict(r) for r in db.q("SELECT i.item_id, i.label, i.value_num, i.display, i.unit, l.status, l.metric_id, l.explanation, l.candidates_json FROM trace_items i JOIN trace_links l ON l.item_id=i.item_id AND l.run_id=i.run_id WHERE i.run_id=? ORDER BY i.rowid", (rid_l,))]
        for it in items:
            it["candidates"] = json.loads(it.pop("candidates_json") or "[]")
        return {**base, "kind": "htm", "items": items, "real": True}
    if path.suffix == ".json":
        return {**base, "kind": "json", "doc": json.loads(path.read_text(encoding="utf-8"))}
    return base


# ------------------------------------------------------------------------------------------------------------------ the ripple and the graph
LANES = {"tidewall": ["Sources", "Ledger", "Trackback", "Close Accountant", "Covenant and Comp", "Reporting Analyst", "People", "Actions"],
         "lyft": ["Sources", "Trackback", "Inside (with Trackback)", "Newsroom (outside)", "Sell-side desk (outside)", "Actions"]}
OWNER_LANE = {"close_accountant": "Close Accountant", "covenant_comp": "Covenant and Comp", "reporting_analyst": "Reporting Analyst", "priya": "People", "sam": "People",
              "newsroom": "Newsroom (outside)", "street_analyst": "Sell-side desk (outside)", "lyft_ir": "Inside (with Trackback)", "lyft_fpa": "Inside (with Trackback)"}


@router.get("/ripple/{exception_id}")
def ripple(exception_id: str, world: str = "tidewall"):
    rid = _latest(world)
    r = db.q1("SELECT result_json FROM ripples WHERE run_id=? AND exception_id=?", (rid, exception_id))
    if r is None:
        raise HTTPException(404, "no ripple for this exception")
    rp = json.loads(r["result_json"])
    waves = _kv(f"waves.{rid}", {})
    w = waves.get(exception_id) or next((v for k, v in waves.items() if exception_id in k.split("+")), None)
    return {**rp, "wave": w, "graph": _graph(rp, w, world)}


def _graph(rp: dict, wave: dict | None, world: str) -> dict:
    """Lanes for React Flow: correction -> ledger slices -> changed metrics -> files and agents' work -> actions. Cleared items are grouped."""
    nodes, edges = [], []
    cid = rp["correction"]["id"]
    nodes.append({"id": cid, "lane": "Sources", "kind": "correction", "label": rp["correction"].get("title", cid), "verdict": "invalidated", "hop": 0})
    for s in rp.get("ledger_slices", [])[:4]:
        nodes.append({"id": f"slice:{s}", "lane": "Ledger" if world == "tidewall" else "Sources", "kind": "ledger", "label": s, "verdict": "invalidated", "hop": 1})
        edges.append({"source": cid, "target": f"slice:{s}"})
    changed = {m: v for m, v in rp["metrics"].items() if v["changed"]}
    key_metrics = sorted(changed, key=lambda m: (changed[m]["depth"], m))
    shown = [m for m in key_metrics if not m.startswith("acct.") and "@2026Q2" not in m and "deferred" not in m][:9]
    for m in shown:
        v = changed[m]
        nodes.append({"id": f"M:{m}", "lane": "Trackback", "kind": "metric", "label": v["label"], "old": v["old"], "new": v["new"], "unit": v["unit"],
                      "verdict": "invalidated", "reason": v["reason"], "hop": 2 + v["depth"]})
        for s in rp.get("ledger_slices", [])[:4]:
            if v["seed"]:
                edges.append({"source": f"slice:{s}", "target": f"M:{m}"})
    pings = {i["artifact_id"]: i for p in (wave or {}).get("pings", []) for i in p["items"]}
    for iid, v in rp["items"].items():
        if v["verdict"] != "invalidated":
            continue
        owner = v.get("owner") or "People"
        nodes.append({"id": iid, "lane": OWNER_LANE.get(owner, "People"), "kind": "item", "label": f"{v['file'].split('_')[0]} {v['locator']}: {v['label']}"[:80], "old": v.get("old"), "new": v.get("new"),
                      "verdict": "invalidated", "reason": v["reason"], "owner": owner, "file": v["file"], "ping": pings.get(iid, {}).get("verdict"), "hop": 6})
        if v.get("metric") and f"M:{v['metric']}" in {n["id"] for n in nodes}:
            edges.append({"source": f"M:{v['metric']}", "target": iid})
    for nid, v in rp.get("agent_nodes", {}).items():
        nodes.append({"id": nid, "lane": OWNER_LANE.get(v["owner"], v["owner"]), "kind": "agent", "label": (v.get("content") or nid)[:80], "verdict": v["verdict"], "reason": v["reason"],
                      "owner": v["owner"], "skipped": v.get("skipped"), "ping": pings.get(nid, {}).get("verdict"), "hop": 6})
        for c in v.get("consumes", []):
            if f"M:{c}" in {n["id"] for n in nodes}:
                edges.append({"source": f"M:{c}", "target": nid})
    for aid, v in rp["actions"].items():
        nodes.append({"id": aid, "lane": "Actions", "kind": "action", "label": v["label"], "verdict": v["verdict"], "reason": v["reason"], "old": v.get("old"), "new": v.get("new"),
                      "compensation": v.get("compensation"), "owner": v.get("owner"), "ping": pings.get(aid, {}).get("verdict"), "hop": 7})
        if f"M:{v['metric']}" in {n["id"] for n in nodes}:
            edges.append({"source": f"M:{v['metric']}", "target": aid})
    cleared = [v for v in rp["items"].values() if v["verdict"] == "still_valid"]
    by_reason: dict[str, int] = {}
    for v in cleared:
        by_reason[v["reason"]] = by_reason.get(v["reason"], 0) + 1
    return {"lanes": LANES.get(world, LANES["tidewall"]), "nodes": nodes, "edges": edges, "cleared_by_reason": by_reason}


@router.get("/pings")
def pings(world: str = "tidewall"):
    rid = _latest(world)
    waves = _kv(f"waves.{rid}", {})
    rows = [dict(r) | {"payload": None} for r in db.q("SELECT ping_id, correction_id, wave_id, target_agent_id, state, refusal_code, hop, created_at, responded_at FROM pings ORDER BY created_at, rowid")]
    return {"waves": waves, "pings": rows}


@router.get("/ping/{ping_id}")
def ping(ping_id: str):
    from .core import pings as P
    try:
        return P.status(ping_id)
    except KeyError:
        raise HTTPException(404, "no such ping")


@router.get("/built/{artifact_id}")
def built(artifact_id: str, world: str = "tidewall"):
    """Who built this, on top of whom: the upstream chain across lanes back to sources (agent, run, version, pin), or the traced link for people's lines."""
    out, seen = [], set()
    rid = _latest(world)

    def walk(aid: str, ver: int | None, depth: int):
        if depth > 8 or (aid, ver) in seen:
            return
        seen.add((aid, ver))
        a = db.q1("SELECT * FROM artifacts WHERE artifact_id=? AND version=?", (aid, ver)) if ver else db.artifact_current(aid)
        if a is None:
            tl = db.q1("SELECT i.*, l.status, l.metric_id, l.snapshot, l.confidence, l.explanation FROM trace_items i JOIN trace_links l ON l.item_id=i.item_id AND l.run_id=i.run_id WHERE i.run_id=? AND i.item_id=?", (rid, aid))
            if tl:
                out.append({"depth": depth, "id": aid, "kind": "traced line", "who": tl["owner"], "label": f"{tl['file']} {tl['locator']}: {tl['label']}", "how": f"matched by the tracer ({tl['status']}, confidence {tl['confidence']})",
                            "explanation": tl["explanation"]})
                if tl["metric_id"]:
                    for m in (tl["metric_id"][4:].split("+") if tl["metric_id"].startswith("sum:") else [tl["metric_id"]]):
                        walk(f"M:{m}", None, depth + 1)
            return
        run = db.q1("SELECT * FROM runs WHERE run_id=?", (a["run_id"],))
        out.append({"depth": depth, "id": aid, "version": a["version"], "kind": a["node_kind"] or a["kind"], "who": a["agent_id"], "label": a["content"], "value": a["value_num"],
                    "run": a["run_id"], "pinned_at": a["pinned_at"], "task": run["task"] if run else None, "lot": db.lot_code(aid, a["version"]),
                    "how": "consumption receipt (captured at read)" if a["agent_id"] not in ("trackback",) else "computed by Trackback"})
        sups = G.supporters_of(aid, a["version"])
        for e in sups:
            if e["supporter_kind"] == "artifact":
                walk(e["supporter_id"], int(e["supporter_version"]), depth + 1)
        if aid.startswith("M:") and not sups:
            m = aid[2:]
            out.append({"depth": depth + 1, "id": f"ledger:{m}", "kind": "ledger", "who": "general ledger", "label": "posted entries (Books as of the run's pin)", "how": "the metric's slices"})

    walk(artifact_id, None, 0)
    if world == "tidewall" and out:
        out.append({"depth": max(o["depth"] for o in out) + 1, "id": "billing_release_lines@v1", "kind": "source", "who": "billing system", "label": "Billing line-level export captured Apr 9 18:00 (batch B07 released once)",
                    "how": "system of record for annual-plan releases"})
    return {"artifact_id": artifact_id, "chain": out}


# ------------------------------------------------------------------------------------------------------------------ approvals, people, clock
class Decision(BaseModel):
    person: str
    decision: str
    reason: str = ""


@router.get("/approvals")
def approvals(world: str = "tidewall"):
    from . import workflow
    if world == "lyft":
        rows = []
        for r in db.q("SELECT * FROM actions WHERE world='lyft' AND kind NOT IN ('publish_story','order') ORDER BY created_at, action_id"):
            d = dict(r); d["params"] = json.loads(d.pop("params_json") or "{}"); d["approvals"] = [dict(x) for x in db.q("SELECT * FROM approvals WHERE action_id=? ORDER BY ts", (r["action_id"],))]
            rows.append(d)
        return {"items": rows, "holds": [], "clock": None}
    return {"items": workflow.inbox(), "holds": workflow.holds(), "clock": workflow.clock(),
            "challenges": [dict(r) for r in db.q("SELECT * FROM upstream_challenges ORDER BY created_at")]}


@router.post("/approvals/{action_id}/decision")
def decide(action_id: str, d: Decision, world: str = "tidewall"):
    from . import workflow
    _snapshot_for_undo()
    try:
        if world == "lyft":
            from .world import lyft as LY
            return LY.decide(action_id, d.person, d.decision, d.reason)
        return workflow.decide(action_id, d.person, d.decision, d.reason)
    except AP.SelfApprovalError as e:
        return JSONResponse({"refused": True, "error": f"segregation of duties: {e}", "check": "CHECK (proposer <> approver)"}, status_code=409)
    except KeyError:
        raise HTTPException(404, "no such item")
    except PermissionError as e:
        return JSONResponse({"refused": True, "error": str(e)}, status_code=409)


@router.post("/people/{person}/act")
def person_act(person: str, world: str = "tidewall"):
    _snapshot_for_undo()
    if world == "lyft":
        from .world import lyft as LY
        return {"person": person, "results": LY.simulate(person)}
    from . import workflow
    return {"person": person, "results": workflow.simulate(person)}


class Advance(BaseModel):
    hours: float = 24


@router.post("/clock/advance")
def clock_advance(a: Advance):
    from . import workflow
    _snapshot_for_undo()
    return workflow.advance(a.hours)


# ------------------------------------------------------------------------------------------------------------------ memo, lessons, scorecard
@router.get("/memo/{exception_id}")
def memo_view(exception_id: str, world: str = "tidewall", rebuild: bool = False):
    from . import memo
    if rebuild and world == "tidewall":
        memo.build(_latest("tidewall"), exception_id, C.KEY_DATES["now"])
    r = db.q1("SELECT * FROM memos WHERE exception_id=? ORDER BY version DESC LIMIT 1", (exception_id,))
    if r is None:
        raise HTTPException(404, "no memo")
    return {"exception_id": exception_id, "version": r["version"], "markdown": r["markdown"], "html": memo.to_html(r["markdown"])}


@router.get("/memos")
def memos(world: str = "tidewall"):
    return {"memos": [dict(r) for r in db.q("SELECT memo_id, MAX(version) version, exception_id, world FROM memos WHERE world=? GROUP BY memo_id", (world,))]}


@router.get("/lessons")
def lessons_view():
    from . import lessons
    ev = _eval_file()
    return {"lessons": lessons.listing(), "learning": ev.get("learning")}


@router.post("/lessons/{lesson_id}/approve")
def lesson_approve(lesson_id: str, d: Decision):
    from . import lessons
    _snapshot_for_undo()
    if d.person == "close_accountant":
        return JSONResponse({"refused": True, "error": "the agent cannot approve its own lesson"}, status_code=409)
    return lessons.approve(lesson_id, d.person, C.KEY_DATES["now"])


@router.post("/lessons/{lesson_id}/gate")
def lesson_gate(lesson_id: str):
    """Run the regression gate live: the whole case library without and with the lesson, on top of the lessons already active."""
    from . import lessons
    return lessons.gate_live(lesson_id)


@router.post("/review/rerun")
def review_rerun():
    """The same review under the current playbook (v1 plus approved lessons), next to the shipped v1 run. Nothing is re-sent or re-proposed."""
    from . import review
    now = review.evaluate_only()
    rid = _latest("tidewall")
    before = {}
    for r in db.q("SELECT exception_id, claims_json FROM evaluations WHERE run_id=? AND playbook_version NOT LIKE '%+jade'", (rid,)):
        ev = json.loads(r["claims_json"])
        js = [dict(x) for x in db.q("SELECT claim_key, verdict, why FROM jade_reviews WHERE eval_id=? AND verdict IN ('overrule','flag','escalate')", (f"EV-{rid}-{r['exception_id']}",))]
        before[r["exception_id"]] = {"title": ev["title"], "classification": ev["q3_classification"]["value"], "material": ev["q5_matters"]["material"], "posting": ev["q7_fix"].get("posting_period"),
                                     "treatment": ev["q7_fix"].get("treatment"), "approvers": ev["q7_fix"].get("approvers"), "jade": [{"claim": x["claim_key"], "verdict": x["verdict"], "why": x["why"]} for x in js]}
    return {"before": {"playbook": "v1", "exceptions": before}, "now": now}


@router.get("/scorecard")
def scorecard():
    ev = _eval_file()
    if not ev:
        raise HTTPException(404, "run python tasks.py eval")
    import hashlib
    return {**ev, "key_sha256_now": hashlib.sha256(C.ANSWER_KEY_V3.read_bytes()).hexdigest() if C.ANSWER_KEY_V3.exists() else None}


# ------------------------------------------------------------------------------------------------------------------ replay (time travel)
@router.get("/memento/timemap/{source}")
def timemap(source: str):
    from .engine import replay as R
    try:
        return {"source": source, "title": db.q1("SELECT title FROM sources WHERE source_id=?", (source,))["title"], "versions": [{k: v[k] for k in ("version_id", "captured_at")} | {"note": v["body"].get("note")} for v in R.timemap(source)]}
    except (KeyError, TypeError):
        raise HTTPException(404, "no such source")


@router.get("/memento/timegate/{source}")
def timegate(source: str, at: str):
    from .engine import replay as R
    try:
        v = R.timegate(source, at, "ui-replay")
        return {"source": source, "pinned_at": db.norm_ts(at), "version": {k: v[k] for k in ("version_id", "captured_at")}}
    except R.Refusal as e:
        return JSONResponse(e.problem(), status_code=e.status, media_type="application/problem+json")


class Attempt(BaseModel):
    source: str
    version_id: str | None = None
    pin: str = "2026-04-09T18:00:00-07:00"


@router.post("/replay/attempt")
def replay_attempt(a: Attempt):
    """Door 3, 'Try to make it peek': request a later version under an earlier pin. Refused with both timestamps."""
    from .engine import replay as R
    vs = R.timemap(a.source)
    target = next((v for v in vs if v["version_id"] == a.version_id), vs[-1])
    try:
        R.fetch(target["version_id"], a.pin, "ui-cheat")
        return {"served": True, "version_id": target["version_id"]}
    except R.Refusal as e:
        return {"served": False, "refused": True, "problem": e.problem(), "status": e.status}


@router.get("/ledger")
def ledger(as_of: str | None = None, world: str = "tidewall"):
    from .engine.books import Books
    from .world.build import load_entries
    at = db.norm_ts(as_of) if as_of else None
    ents = [e for e in load_entries() if at is None or db.norm_ts(e["posted_at"]) <= at]
    tb = Books(ents).trial_balance("2026-06")
    return {"as_of": as_of, "entries": len(ents), "trial_balance": tb}


ASOF_METRICS = {"tidewall": [("ttm_adj_ebitda@2026Q1", "TTM Adjusted EBITDA (Q1)"), ("leverage@2026Q1", "Net leverage (Q1)"), ("covenant_status@2026Q1", "Covenant status (Q1)"),
                             ("revenue@2026Q1", "Revenue, Q1 2026"), ("cash@2026Q1", "Cash, March 31"), ("acct.4010@2026Q1", "Annual-plan revenue (4010), Q1")],
                "lyft": [("release.margin_expansion_bps", "Margin expansion in the public release", "bps"), ("lyft.margin_expansion_bps", "Margin expansion in the guidance model", "bps"),
                         ("street.implied_fy24_adj_ebitda", "Sell-side implied FY24 Adjusted EBITDA")]}


@router.get("/asof")
def asof(at: str, world: str = "tidewall"):
    """The books as the system held them at a moment (posted_at at or before it), for the time-travel scrubber."""
    from .engine.books import Books
    t = db.norm_ts(at)
    if world == "lyft":
        from .world import lyft as LY
        ents = [e for e in LY.entries() if db.norm_ts(e["posted_at"]) <= t]
        vals = LY.library().evaluate(Books(ents))
    else:
        from . import review
        cx = review.ctx(at)
        vals = cx.lib.evaluate(Books(cx.entries))
        ents = cx.entries
    return {"at": at, "entries": len(ents), "metrics": [{"id": m[0], "label": m[1], "unit": m[2] if len(m) > 2 else None, "value": vals.get(m[0])} for m in ASOF_METRICS[world]]}


@router.get("/decision/{action_id}")
def decision_replay(action_id: str):
    """Door 2: 'Why did you do this on that date?' The action's inputs exactly as of its pin."""
    a = db.q1("SELECT * FROM actions WHERE action_id=?", (action_id,))
    if a is None:
        raise HTTPException(404, "no such action")
    chain = built(a["artifact_id"], a["world"] or "tidewall")["chain"]
    return {"action": dict(a), "executed_at": a["executed_at"], "inputs": chain,
            "approvals": [dict(r) for r in db.q("SELECT * FROM approvals WHERE action_id=? ORDER BY ts", (action_id,))]}


# ------------------------------------------------------------------------------------------------------------------ the judge's own correction
class CorrectionReq(BaseModel):
    free_text: str | None = None
    kind: str | None = None                 # ledger | billing_duplicate | invoice
    account: str | None = None
    period: str | None = None
    entity: str | None = None
    amount: float | None = None             # dollars; positive increases the account's debit
    batch: str | None = None
    knowable: str = "before"                # before | after the close


def _parse(text: str) -> dict:
    t = text.lower()
    m = re.search(r"(?<![a-z\d])\$?\s*(\d[\d,]*(?:\.\d+)?)\s*(k|m)?\b", t)
    amt = float(m.group(1).replace(",", "")) * ({"k": 1e3, "m": 1e6}.get(m.group(2) or "", 1)) if m else None
    months = {"january": "01", "february": "02", "march": "03", "april": "04", "may": "05", "june": "06", "jan": "01", "feb": "02", "mar": "03", "apr": "04", "jun": "06"}
    mo = next((v for k, v in months.items() if re.search(rf"\b{k}\b", t)), "03")
    period = f"2026-{mo}"
    acct = next((a for a, (n, _t) in C.ACCOUNTS.items() if n.split(" (")[0].lower() in t), None)
    for word, a in (("legal", "6400"), ("utilit", "6200"), ("route", "6300"), ("payroll", "6000"), ("rent", "6100"), ("dues", "4000"), ("annual", "4010"), ("day pass", "4100"), ("wellness", "4200"), ("wall", "4300"), ("bank fee", "6500")):
        if word in t:
            acct = a
    batch = re.search(r"\bb0?(\d)\b", t)
    kind = "billing_duplicate" if (("duplicat" in t or "twice" in t) and ("batch" in t or batch or "import" in t)) else ("invoice" if "invoice" in t else "ledger")
    knowable = "after" if re.search(r"after the close|arrived (in|on) (may|june|july)|received (in|on) (may|june|july)", t) else "before"
    region = next((r for r, n in C.REGIONS.items() if n.lower() in t), None)
    entity = next((g for g, (r, c) in C.GYM_IDS.items() if c.lower() in t), None) or (next((g for g, (r, c) in C.GYM_IDS.items() if r == region), None) if region else "HQ")
    return {"kind": kind, "amount": amt, "period": period, "account": acct or "6700", "entity": entity, "batch": f"B0{batch.group(1)}" if batch else None, "knowable": knowable}


@router.post("/corrections/parse")
def parse_correction(req: CorrectionReq):
    return {"parsed": _parse(req.free_text or "")}


@router.post("/corrections")
def correction(req: CorrectionReq):
    """'Change any number': build the correcting lines, run the ripple, materiality, the agent and Jade on it, live. Nothing is posted."""
    from . import jade, review as RV
    from .agent import agent as AG, playbook as PB
    from .engine import materiality as MAT, replay as R
    p = _parse(req.free_text) if req.free_text and not req.kind else req.model_dump()
    pin = C.KEY_DATES["now"]
    if p["kind"] == "billing_duplicate":
        from .eval.cases import dup_import
        from .world import generator as GEN
        w = GEN.build()
        _mut, lines, total, by_region = dup_import(p.get("batch") or "B03", p["period"], w.members, f"{p['period']}-28T23:00:00-07:00")
        exc = {"kind": "duplicate_import", "title": f"Judge: batch {p.get('batch') or 'B03'} {p['period']} imported twice", "amount": total, "by_region": by_region, "batch": p.get("batch") or "B03", "period": p["period"],
               "duplicate_import": "JUDGE", "evidence": [{"what": "GL import", "import_id": "JUDGE", "posted_at": f"{p['period']}-28T23:00:00-07:00"}]}
    else:
        amt = int(round((p.get("amount") or 0) * 100))
        if not amt:
            raise HTTPException(400, "give an amount (for example: 'March legal is $90k higher')")
        lines = [{"entity": p.get("entity") or "HQ", "account": p["account"], "amount_cents": amt, "orig_period": p["period"]},
                 {"entity": "HQ", "account": "2010" if C.ACCOUNTS[p["account"]][1] == "expense" else "2300", "amount_cents": -amt, "orig_period": p["period"]}]
        received = "2026-04-02T09:00:00-07:00" if p.get("knowable", "before") == "before" else "2026-05-20T09:00:00-07:00"
        exc = {"kind": "unaccrued_invoice" if C.ACCOUNTS[p["account"]][1] == "expense" else "misclassification", "title": f"Judge: {C.ACCOUNTS[p['account']][0]} {p['period']} {'+' if amt > 0 else ''}{amt / 100:,.0f}",
               "amount": abs(amt) / 100, "evidence": [], "knowable": {"at": received, "fact": f"the fact first existed {received[:10]}"}}
    exc.update(exception_id="JUDGE", recon="judge", correction={"lines": lines, "orig_period": p["period"]}, system_of_record="judge", trust={"trusted": "the judge", "over": "the books", "why": "a judge's correction"})
    links = RV.trace(pin, store=False)
    from .world import history
    rp = RV.ripple_for(exc, links, pin, store=False, agent_nodes=history.agent_nodes())
    cls = AG.classify(exc, PB.apply(PB.V1, ["classify_by_knowable_at"]), AG.knowable(exc))["value"]
    mat = MAT.evaluate([{**exc, "classification": cls}], R.rows("misstatement_log", pin), {"JUDGE": rp})
    ev = AG.evaluate(exc, rp, mat, PB.current(), pin)
    jf = jade.review(ev, exc, mat, pin)
    fin = jade.apply(ev, jf)
    from . import hub
    agents, people, skipped = hub.bundles(rp)
    return {"parsed": p, "correction": exc, "ripple": rp, "graph": _graph(rp, None, "tidewall"), "evaluation": ev, "jade": jf, "final": fin, "materiality": mat["per_item"].get("JUDGE"),
            "aggregate": mat["aggregate"], "would_ping": {a: [n["item_id"] for n in v] for a, v in agents.items()}, "would_notify": {p_: len(v) for p_, v in people.items()}, "skipped": [s["node_id"] for s in skipped]}


# ------------------------------------------------------------------------------------------------------------------ iteration-2 live actions
@router.post("/s1")
def s1_self_revision():
    from . import hub
    _snapshot_for_undo()
    return hub.revise("CA.june_accrual.utilities", agent="close_accountant", new_value=195_000.0, reason="Bayline's June meter read came in higher than the estimate", ts="2026-07-08T09:20:00-07:00")


@router.get("/events")
def events(world: str | None = None, limit: int = 200):
    rows = db.q("SELECT * FROM events ORDER BY event_id DESC LIMIT ?", (limit,))
    return {"events": [dict(r) | {"detail": json.loads(r["detail_json"]) if r["detail_json"] else None} for r in rows]}


@router.get("/access")
def access():
    return {"reads": ["billing release lines", "GL detail and import log", "bank statement and cash book", "AR and AP registers", "accruals and the misstatement log", "contracts and policies", "people's files (covenant workbook, board deck, certificate)"],
            "drafts_only": ["journal entries (posted only after the Controller approves)", "corrected certificate and lender notice (CFO signs)", "bonus offsets (HR approves)", "vendor credits and holds (AP approves)", "the error memo"],
            "never": ["post to the ledger", "approve its own work (CHECK proposer <> approver in the schema)", "send anything to the lender", "pay or move money", "edit a locked period", "change a source system"],
            "evidence": "every claim cites the records it rests on; every read is logged with its pin; refusals included"}


@router.post("/reset")
def reset():
    db.restore_snapshot(C.SNAPSHOT_PATH)
    if UNDO_DIR.exists():
        shutil.rmtree(UNDO_DIR, ignore_errors=True)
    db.reset_connection()
    return {"reset": True}


@router.post("/undo")
def undo():
    snaps = sorted(UNDO_DIR.glob("*.sqlite")) if UNDO_DIR.exists() else []
    if not snaps:
        return {"undone": False}
    db.restore_snapshot(snaps[-1]); snaps[-1].unlink(); db.reset_connection()
    return {"undone": True, "remaining": len(snaps) - 1}


# ------------------------------------------------------------------------------------------------------------------ the Stage bundle
@router.get("/stage")
def stage(world: str = "tidewall"):
    """Everything Stage mode needs, in one call, from the stored run."""
    if world == "lyft":
        return _stage_lyft()
    from . import workflow
    rid = _latest("tidewall")
    exc = exceptions("tidewall")["exceptions"]
    hero = next(e for e in exc if e["kind"] == "duplicate_import")
    ev = evaluation(hero["exception_id"])
    rp = ripple(hero["exception_id"])
    st = state("tidewall")
    rc = recons("tidewall")["summary"]
    lev = rp["metrics"].get("leverage@2026Q1", {})
    head = rp["metrics"].get("covenant_headroom@2026Q1", {})
    links = {l["item_id"]: l for f in lineage("tidewall")["files"].values() for l in f}
    evid = {"billing": next((e for e in hero["evidence"] if e.get("what", "").startswith("billing")), None), "imports": [e for e in hero["evidence"] if e.get("what") == "GL import"],
            "workbook": links.get("WB.Covenant.B7"), "deck": links.get("DECK.s2.l2"), "deck_covenant": links.get("DECK.s3.l1"), "certificate": links.get("CERT.l1"), "by_region": hero.get("by_region")}
    mat = materiality("tidewall")
    ap = workflow.inbox()
    return {"world": "tidewall", "run_id": rid, "now": C.KEY_DATES["now"], "company": C.COMPANY_V3, "title": {"leverage": lev.get("old"), "max": C.COVENANT["max_leverage"], "headroom": head.get("old"),
            "revenue_ttm": links.get("WB.Covenant.B6", {}).get("value_num"), "loan": C.COVENANT["facility"], "gyms": 48},
            "recons": rc, "exceptions": exc, "hero": hero, "evidence": evid, "ripple": {k: rp[k] for k in ("counts", "order", "items", "actions", "agent_nodes", "metrics", "wave", "graph", "seconds")},
            "evaluation": ev, "materiality": {k: mat.get(k) for k in ("aggregate", "rows", "thresholds", "per_item")}, "approvals": [a for a in ap if a.get("exception_id") == hero["exception_id"] or a["kind"] in ("notice", "upstream_challenge")],
            "holds": workflow.holds(), "clock": workflow.clock(), "scorecard": (_eval_file() or {}).get("headline"), "configs": (_eval_file() or {}).get("configs"),
            "replay": {"source": "billing_release_lines", "pin": C.KEY_DATES["draft_tb"], "close": C.KEY_DATES["q1_close"]}, "seat_line": SEAT_LINE, "tagline": TAGLINE, "mode": st["mode"], "cache_rows": st["cache_rows"]}


SEAT_LINE = "Trackback does the close accountant's error review: when sources disagree, it finds what broke, sizes it, drafts the fix and the memo, and never posts anything without the controller."
TAGLINE = "Hyphenate traces any number back. Trackback traces any correction forward."


def _stage_lyft() -> dict:
    from .world import lyft as LY
    rid = _latest("lyft")
    exc = exceptions("lyft")["exceptions"]
    hero = exc[0] if exc else None
    ev = evaluation(hero["exception_id"], "lyft") if hero else None
    rp = ripple(hero["exception_id"], "lyft") if hero else None
    fence = db.q1("SELECT * FROM events WHERE kind='fence_block' AND summary LIKE '%LY.ir%' ORDER BY event_id DESC LIMIT 1")
    lin = lineage("lyft")
    return {"world": "lyft", "run_id": rid, "tape": LY.TAPE, "times": LY.T, "hero": hero, "evaluation": ev, "ripple": {k: rp[k] for k in ("counts", "items", "actions", "agent_nodes", "metrics", "wave", "graph")} if rp else None,
            "trace": lin, "fence": json.loads(fence["detail_json"]) if fence else None, "approvals": approvals("lyft")["items"], "scorecard": (_eval_file() or {}).get("lyft"),
            "real_vs_reconstructed": {"real": ["Exhibit 99.1 as filed at 16:06 ET (500 basis points)", "the CFO's correction on the 4:30 pm ET call (50)", "Business Wire's CORRECTING and REPLACING release", "the tape (up as much as 67% after hours; closed +35.1% next day)"],
                                      "reconstructed": ["the close pack and the FP&A guidance model (numbers from the release)", "the inside agents and the fence", "the newsroom and the sell-side desk"]}}
